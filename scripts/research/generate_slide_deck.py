#!/usr/bin/env python3
"""
Generate the Dashbud general audience presentation as a PPTX file.
Based on: wiki/dashbud/sales/general-audience-1hr-outline.md
"""
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pathlib import Path

OUTPUT = Path("/Users/arthur/www/dashbud/content/sales/presentations/dashbud-general-audience.pptx")

# Brand colors
TEAL = RGBColor(0x00, 0x7B, 0xA1)
TEAL_DARK = RGBColor(0x00, 0x54, 0x6E)
TEAL_LIGHT = RGBColor(0xE0, 0xF4, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
BLACK = RGBColor(0x11, 0x18, 0x27)
GRAY = RGBColor(0x6B, 0x72, 0x80)
GRAY_LIGHT = RGBColor(0xF3, 0xF4, 0xF6)

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)


def add_slide(title_text=None, bg_color=None):
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout
    if bg_color:
        bg = slide.background
        fill = bg.fill
        fill.solid()
        fill.fore_color.rgb = bg_color
    return slide


def add_text_box(slide, left, top, width, height, text, font_size=18, color=BLACK, bold=False, alignment=PP_ALIGN.LEFT, font_name="Calibri"):
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.name = font_name
    p.alignment = alignment
    return tf


def add_bullet_list(slide, left, top, width, height, items, font_size=16, color=BLACK, spacing=Pt(8), bulleted=True):
    from pptx.oxml.ns import qn
    txBox = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(font_size)
        p.font.color.rgb = color
        p.font.name = "Calibri"
        p.space_after = spacing
        p.level = 0
        if bulleted:
            pPr = p._pPr
            if pPr is None:
                pPr = p._p.get_or_add_pPr()
            buChar = pPr.makeelement(qn('a:buChar'), {'char': '•'})
            # Remove any existing bullet settings
            for child in list(pPr):
                if 'buNone' in child.tag or 'buChar' in child.tag:
                    pPr.remove(child)
            pPr.append(buChar)
    return tf


def section_header(title, subtitle=None):
    slide = add_slide(bg_color=TEAL)
    add_text_box(slide, 1.5, 2.5, 10, 1.5, title, font_size=40, color=WHITE, bold=True)
    if subtitle:
        add_text_box(slide, 1.5, 4.2, 10, 1, subtitle, font_size=22, color=TEAL_LIGHT)
    return slide


# ─── SLIDE 1: Title ───
slide = add_slide(bg_color=TEAL_DARK)
add_text_box(slide, 1.5, 2, 10, 1.5, "Dashbud", font_size=54, color=WHITE, bold=True)
add_text_box(slide, 1.5, 3.5, 10, 1, "AI-powered analytics for all your data", font_size=28, color=TEAL_LIGHT)
add_text_box(slide, 1.5, 5.5, 10, 0.5, "getdashbud.com", font_size=18, color=RGBColor(0x80, 0xBB, 0xCE))


# ─── SECTION 1: AUDIENCE ENGAGEMENT ───
section_header("Let's start with you", "How do you work with data today?")

# Slide: What tools?
slide = add_slide()
add_text_box(slide, 0.8, 0.5, 11, 0.8, "What data tools do you use?", font_size=32, color=BLACK, bold=True)
add_bullet_list(slide, 0.8, 1.5, 5.5, 5, [
    "📊  Spreadsheets — Excel, Google Sheets",
    "📈  Application dashboards — QuickBooks, Google Analytics, Meta Ads, Shopify",
    "🔬  BI / analysis platforms — Tableau, Power BI, Looker",
    "🤖  AI-based workflows — Claude, ChatGPT, Gemini",
    "💻  Code — Python, Jupyter Notebooks",
], font_size=20, spacing=Pt(14))

# Slide: What for + how often
slide = add_slide()
add_text_box(slide, 0.8, 0.5, 5.5, 0.8, "What do you track?", font_size=28, color=BLACK, bold=True)
add_bullet_list(slide, 0.8, 1.5, 5.5, 4, [
    "Sales, revenue, pipeline",
    "Expenses, budgets, profitability",
    "Marketing performance",
    "Operations, inventory, fulfillment",
    "Customer behavior, retention",
    "Compliance, stakeholder reporting",
], font_size=18, spacing=Pt(10))

add_text_box(slide, 7, 0.5, 5.5, 0.8, "How often?", font_size=28, color=BLACK, bold=True)
add_bullet_list(slide, 7, 1.5, 5.5, 4, [
    "Multiple times per day",
    "Daily",
    "Weekly",
    "Monthly / quarterly",
    '"When someone asks me for a number"',
], font_size=18, spacing=Pt(10))


# ─── SECTION 2: THE PROBLEM ───
section_header("The problem", "Data is valuable. Getting answers from it is painful.")

# Slide: Trade-offs table
slide = add_slide()
add_text_box(slide, 0.8, 0.3, 11, 0.8, "Every approach has trade-offs", font_size=32, color=BLACK, bold=True)

# Table
table_data = [
    ["How you do it", "What's good", "What's painful"],
    ["SaaS dashboards\n(QuickBooks, GA, etc.)", "Built in, no setup", "Clunky, siloed, can't cross\ndata sources"],
    ["BI tools\n(Power BI, Tableau)", "Powerful,\nenterprise-grade", "Finicky, hard to learn,\nexpensive, need a data team"],
    ["Spreadsheets\n(Excel, Google Sheets)", "Flexible,\nyou control it", "Fragile, time-consuming\nto update, hard to share"],
    ["AI workflows\n(Claude, ChatGPT)", "Fast for one-offs,\nsurprisingly capable", "Don't scale — you become\nthe infrastructure"],
]

rows, cols = len(table_data), len(table_data[0])
tbl = slide.shapes.add_table(rows, cols, Inches(0.8), Inches(1.3), Inches(11.5), Inches(5.5)).table

for i, row in enumerate(table_data):
    for j, cell_text in enumerate(row):
        cell = tbl.cell(i, j)
        cell.text = cell_text
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.size = Pt(15 if i > 0 else 14)
            paragraph.font.name = "Calibri"
            paragraph.font.bold = (i == 0)
            paragraph.font.color.rgb = WHITE if i == 0 else BLACK
        if i == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = TEAL
        elif i % 2 == 0:
            cell.fill.solid()
            cell.fill.fore_color.rgb = GRAY_LIGHT

# Column widths
tbl.columns[0].width = Inches(3)
tbl.columns[1].width = Inches(4)
tbl.columns[2].width = Inches(4.5)


# Slide: AI workflow trap
slide = add_slide()
add_text_box(slide, 0.8, 0.3, 11, 0.8, "The AI workflow trap", font_size=32, color=BLACK, bold=True)
add_text_box(slide, 0.8, 1.2, 10, 1.2,
    "You can absolutely build a reporting workflow in Claude or ChatGPT. "
    "Get it to write a Python script, hook up a cron job, pipe the output to email. It works!",
    font_size=20, color=GRAY)

add_text_box(slide, 0.8, 2.8, 10, 0.6, "The question is: what happens next?", font_size=24, color=BLACK, bold=True)

add_bullet_list(slide, 0.8, 3.6, 10, 3.5, [
    "Who else on your team can modify that script?",
    "What happens when the CEO wants to filter by region themselves?",
    "What happens when you hire someone new who needs different access?",
    "What happens when the source data changes and the script breaks at 2am?",
], font_size=20, spacing=Pt(14))

add_text_box(slide, 0.8, 6.2, 10, 0.8,
    "You can build yourself a report. But now you're the data team, the DevOps team, and the support desk.",
    font_size=20, color=TEAL, bold=True)


# ─── SECTION 3: WHERE DASHBUD FITS ───
section_header("Where Dashbud fits", "AI-powered analytics for all your data")

# Slide: Start small, grow with you
slide = add_slide()
add_text_box(slide, 0.8, 0.3, 11, 0.8, "Start simple. Scale when ready.", font_size=32, color=BLACK, bold=True)

boxes = [
    ("Start with one analysis", "Upload a spreadsheet, ask a question, get a chart.\nFive minutes from signup to your first insight."),
    ("Share your work — free", "Read-only dashboard access is free. Your stakeholders\ndon't need to learn the tool — they just see the results."),
    ("Scale with your team", "Add data sources. Build workspaces for different teams.\nClone a data model without starting over.\nEveryone gets the right access."),
    ("Non-technical people get real answers", "Dashbud's semantic model means the AI understands your\nbusiness context. Anyone can ask questions in plain English\nand get accurate, trustworthy answers."),
]

for i, (title, desc) in enumerate(boxes):
    y = 1.3 + i * 1.5
    add_text_box(slide, 0.8, y, 4, 0.4, title, font_size=20, color=TEAL, bold=True)
    add_text_box(slide, 0.8, y + 0.4, 11, 0.9, desc, font_size=16, color=GRAY)


# Slide: What you're NOT giving up
slide = add_slide()
add_text_box(slide, 0.8, 0.3, 11, 0.8, "What you're NOT giving up", font_size=32, color=BLACK, bold=True)
add_bullet_list(slide, 0.8, 1.5, 10, 5, [
    "❌  You're NOT replacing your ERP, CRM, or any existing system",
    "❌  You're NOT learning SQL or a new query language",
    "❌  You're NOT hiring a data team",
    "❌  You're NOT building a data warehouse",
    "❌  You're NOT committing to a six-month implementation",
], font_size=22, spacing=Pt(18))

add_text_box(slide, 0.8, 5.5, 10, 0.8,
    "Most teams are up and running in days.",
    font_size=24, color=TEAL, bold=True)


# ─── SECTION 4: HOW IT WORKS ───
section_header("How it works", "Connect → Model → Explore → Share")

steps = [
    ("1. Connect your data", "Upload spreadsheets, connect directly to databases, or sync from cloud sources.\nSmart data loading cleans messy formats automatically."),
    ("2. Model through conversation", "Tell Dashbud what your data means — in plain language.\nThe AI asks questions, you answer. No code, no config files."),
    ("3. Ask questions, get answers", "Type a question in the Data Explorer. Get tables, charts, and insights.\nSmart formatting. Parametrized controls. Same answer every time."),
    ("4. Share with your team", "Save to dashboards. Share with stakeholders.\nEveryone gets the right access — from analysts to executives."),
]

for title, desc in steps:
    slide = add_slide()
    add_text_box(slide, 0.8, 0.5, 11, 0.8, title, font_size=36, color=TEAL, bold=True)
    add_text_box(slide, 0.8, 1.8, 10, 2, desc, font_size=22, color=GRAY)
    add_text_box(slide, 0.8, 5.5, 10, 1, "[PRODUCT SCREENSHOT]", font_size=18, color=RGBColor(0xCC, 0xCC, 0xCC), alignment=PP_ALIGN.CENTER)


# ─── SECTION 5: PRIVACY & SECURITY ───
slide = add_slide()
add_text_box(slide, 0.8, 0.3, 11, 0.8, "Built for trust. Designed for privacy.", font_size=32, color=BLACK, bold=True)

privacy_points = [
    ("🔒  AI never sees your raw data", "The AI only sees field names and schema — never the actual values in your database."),
    ("📐  Deterministic reports", "The AI writes a SQL query once. After that, it's a traditional report.\nSame question, same answer, every time."),
    ("📋  Full audit trail", "Every data import is tracked and can be rolled back.\nYour data stays in your database or in Dashbud's managed infrastructure."),
    ("🚫  No third-party sharing", "We don't share your data with third parties.\nWe don't use it for AI training."),
]

for i, (title, desc) in enumerate(privacy_points):
    y = 1.3 + i * 1.5
    add_text_box(slide, 0.8, y, 10, 0.4, title, font_size=20, color=BLACK, bold=True)
    add_text_box(slide, 1.2, y + 0.45, 10, 0.9, desc, font_size=16, color=GRAY)


# ─── DEMO PLACEHOLDER ───
section_header("Live demo", "Let's see it in action")

slide = add_slide()
add_text_box(slide, 2, 3, 9, 1.5, "[LIVE DEMO HERE]\nFollow the Demo Playbook", font_size=28, color=GRAY, alignment=PP_ALIGN.CENTER)


# ─── USE CASES ───
section_header("Use cases", "How businesses like yours use Dashbud")

cases = [
    ("Product-based business", "Sales trends by channel, inventory tracking,\ncustomer purchasing behavior, rep performance"),
    ("Service-based business", "Profitability by client, utilization,\nmarketing ROI, monthly stakeholder reports"),
    ("Healthcare / regulated", "Operational reports from EMR data\nwithout exposing sensitive data to AI platforms"),
]

slide = add_slide()
add_text_box(slide, 0.8, 0.3, 11, 0.8, "Who is Dashbud for?", font_size=32, color=BLACK, bold=True)

for i, (title, desc) in enumerate(cases):
    x = 0.8 + i * 4.2
    # Card background
    shape = slide.shapes.add_shape(
        1,  # rectangle
        Inches(x), Inches(1.5), Inches(3.8), Inches(4.5)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = GRAY_LIGHT
    shape.line.fill.background()

    add_text_box(slide, x + 0.3, 1.8, 3.2, 0.6, title, font_size=20, color=TEAL, bold=True)
    add_text_box(slide, x + 0.3, 2.5, 3.2, 3, desc, font_size=16, color=GRAY)


# ─── Q&A ───
section_header("Q&A", "What questions do you have?")


# ─── CLOSE ───
slide = add_slide(bg_color=TEAL_DARK)
add_text_box(slide, 1.5, 1.5, 10, 1.2, "Get started today", font_size=44, color=WHITE, bold=True)
add_text_box(slide, 1.5, 3, 10, 1, "getdashbud.com", font_size=28, color=TEAL_LIGHT)
add_text_box(slide, 1.5, 4.2, 10, 2,
    "Free trial: getdashbud.com/signup\n"
    "Book a demo: calendly.com/arthur-evolytix/dashbud-demo-and-setup\n"
    "Contact: hello@mail.getdashbud.com",
    font_size=18, color=RGBColor(0x80, 0xBB, 0xCE))

add_text_box(slide, 1.5, 6.2, 10, 0.8,
    '"If you\'ve got data and you\'re spending too much time wrangling it into answers — that\'s what we built Dashbud for."',
    font_size=16, color=WHITE)


# ─── SAVE ───
OUTPUT.parent.mkdir(parents=True, exist_ok=True)
prs.save(str(OUTPUT))
print(f"Saved {OUTPUT} ({len(prs.slides)} slides)")
